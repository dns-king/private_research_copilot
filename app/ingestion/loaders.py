from __future__ import annotations

import hashlib
import mimetypes
from dataclasses import dataclass
from pathlib import Path
from typing import Any


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown", ".html", ".htm"}


@dataclass
class LoadedDocument:
    path: Path
    text: str
    metadata: dict[str, Any]
    content_hash: str
    mime_type: str | None
    size_bytes: int


def discover_files(path: str | Path, recursive: bool = True) -> list[Path]:
    source = Path(path).expanduser().resolve()
    if source.is_file():
        return [source] if source.suffix.lower() in SUPPORTED_EXTENSIONS else []
    if not source.exists():
        raise FileNotFoundError(f"Path does not exist: {source}")
    iterator = source.rglob("*") if recursive else source.glob("*")
    return sorted(
        file for file in iterator if file.is_file() and file.suffix.lower() in SUPPORTED_EXTENSIONS
    )


def load_document(path: str | Path) -> LoadedDocument:
    file_path = Path(path).expanduser().resolve()
    suffix = file_path.suffix.lower()
    raw = file_path.read_bytes()
    content_hash = hashlib.sha256(raw).hexdigest()
    mime_type, _ = mimetypes.guess_type(file_path)
    metadata: dict[str, Any] = {
        "extension": suffix,
        "file_name": file_path.name,
        "parent": str(file_path.parent),
    }

    if suffix == ".pdf":
        text, extra = _load_pdf(file_path)
        metadata.update(extra)
    elif suffix == ".docx":
        text, extra = _load_docx(file_path)
        metadata.update(extra)
    elif suffix in {".html", ".htm"}:
        text, extra = _load_html(raw)
        metadata.update(extra)
    else:
        text = _decode_text(raw)
        if suffix in {".md", ".markdown"}:
            metadata["title"] = _markdown_title(text)

    return LoadedDocument(
        path=file_path,
        text=_normalize_text(text),
        metadata=metadata,
        content_hash=content_hash,
        mime_type=mime_type,
        size_bytes=file_path.stat().st_size,
    )


def _load_pdf(path: Path) -> tuple[str, dict[str, Any]]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("Install pypdf to ingest PDF files") from exc

    reader = PdfReader(str(path))
    pages = []
    for index, page in enumerate(reader.pages):
        page_text = page.extract_text() or ""
        if page_text.strip():
            pages.append(f"\n\n[page {index + 1}]\n{page_text}")
    metadata = {
        "page_count": len(reader.pages),
        "pdf_metadata": {
            str(key).strip("/"): str(value) for key, value in (reader.metadata or {}).items()
        },
    }
    return "\n".join(pages), metadata


def _load_docx(path: Path) -> tuple[str, dict[str, Any]]:
    try:
        import docx
    except ImportError as exc:
        raise RuntimeError("Install python-docx to ingest DOCX files") from exc

    document = docx.Document(str(path))
    blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    props = document.core_properties
    metadata = {
        "title": props.title or None,
        "author": props.author or None,
        "created": props.created.isoformat() if props.created else None,
        "modified": props.modified.isoformat() if props.modified else None,
    }
    return "\n\n".join(blocks), {k: v for k, v in metadata.items() if v}


def _load_html(raw: bytes) -> tuple[str, dict[str, Any]]:
    try:
        from bs4 import BeautifulSoup
    except ImportError as exc:
        raise RuntimeError("Install beautifulsoup4 to ingest HTML files") from exc

    soup = BeautifulSoup(raw, "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    title = soup.title.string.strip() if soup.title and soup.title.string else None
    text = soup.get_text("\n")
    return text, {"title": title} if title else {}


def _decode_text(raw: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


def _normalize_text(text: str) -> str:
    lines = [line.strip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    normalized = "\n".join(line for line in lines if line)
    return normalized.strip()


def _markdown_title(text: str) -> str | None:
    for line in text.splitlines():
        line = line.strip()
        if line.startswith("# "):
            return line[2:].strip()
    return None

